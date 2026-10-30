#!/usr/bin/env python3
"""Execute the pinned Python readers on independently written document fixtures.

Development only: python-docx, openpyxl, python-pptx, xlrd, xlwt, olefile, striprtf.
Production and the ordinary Go test suite require none of these packages.
"""
import ast
import base64
import hashlib
import io
import json
from pathlib import Path
import re
import struct
import subprocess
import sys
from types import ModuleType, SimpleNamespace
import zipfile

BASELINE = '6605cb3474cf0c2894080af216d5fed2a81fc997'
root = Path(sys.argv[1]).resolve()
destination = Path(__file__).resolve().parents[2] / 'server/internal/app'
sources = {}
def source(path):
    raw = subprocess.check_output(['git', 'show', f'{BASELINE}:{path}'], cwd=root)
    sources[path] = hashlib.sha256(raw).hexdigest()
    return raw

config = ModuleType('app.config')
config.settings = SimpleNamespace()
sys.modules['app.config'] = config
module = ModuleType('upstream_text_extraction')
sys.modules[module.__name__] = module
exec(compile(source('backend/app/services/text_extraction.py'), 'upstream_text_extraction.py', 'exec'), module.__dict__)
service = module.TextExtractionService()

def archive(files):
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in sorted(files.items()):
            z.writestr(zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0)), data, compress_type=zipfile.ZIP_DEFLATED)
    return output.getvalue()

def normalize_zip(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        files = {name: z.read(name) for name in z.namelist()}
        if 'docProps/core.xml' in files:
            files['docProps/core.xml'] = re.sub(rb'(<dcterms:(?:created|modified)[^>]*>)[^<]*', rb'\g<1>2000-01-01T00:00:00Z', files['docProps/core.xml'])
        return archive(files)

def compound(streams):
    """Small CFB fixture writer: regular streams, one FAT sector, no mini-stream.

    The pinned Python olefile reader independently validates these containers.
    """
    END, FREE = 0xFFFFFFFE, 0xFFFFFFFF
    entries = [('Root Entry', 5, '', b'')]
    for name in sorted(streams):
        parent, _, leaf = name.rpartition('/')
        if parent and not any(e[0] == parent and e[1] == 1 for e in entries):
            entries.append((parent, 1, '', b''))
        entries.append((leaf, 2, parent, streams[name]))
    directory_sectors = (len(entries)+3)//4
    sectors = [bytearray(512) for _ in range(directory_sectors)]
    fat = [FREE]*128
    for i in range(directory_sectors): fat[i] = i+1 if i+1 < directory_sectors else END
    positions = {}
    for i, (_, kind, _, data) in enumerate(entries):
        if kind != 2: continue
        # Regular streams must be >=4096 bytes. Padding is outside the logical
        # text/piece table; string property readers trim terminal NULs.
        data = data.ljust(max(4096, ((len(data)+511)//512)*512), b'\0')
        first = len(sectors)
        for off in range(0, len(data), 512):
            index = len(sectors); sectors.append(bytearray(data[off:off+512]))
            fat[index] = index+1 if off+512 < len(data) else END
        positions[i] = (first, len(data))
    links = {}
    def tree(ids):
        if not ids: return FREE
        middle = len(ids)//2; node = ids[middle]
        links[node] = (tree(ids[:middle]), tree(ids[middle+1:]))
        return node
    children = {}
    for i, (name, kind, _, _) in enumerate(entries):
        if kind in (1, 5):
            parent = name if kind == 1 else ''
            ids = [j for j, e in enumerate(entries) if j and e[2] == parent]
            ids.sort(key=lambda j: (len(entries[j][0]), entries[j][0].upper()))
            children[i] = tree(ids)
    for i, (name, kind, _, _) in enumerate(entries):
        entry = bytearray(128); encoded = (name+'\0').encode('utf-16le')
        entry[:len(encoded)] = encoded
        struct.pack_into('<HBB', entry, 64, len(encoded), kind, 1)
        struct.pack_into('<III', entry, 68, *links.get(i, (FREE, FREE)), children.get(i, FREE))
        struct.pack_into('<IQ', entry, 116, *positions.get(i, (END, 0)))
        start = (i%4)*128; sectors[i//4][start:start+128] = entry
    fat_sector = len(sectors); fat[fat_sector] = 0xFFFFFFFD
    if fat_sector >= 128: raise ValueError('fixture too large')
    header = bytearray(512); header[:8] = bytes.fromhex('d0cf11e0a1b11ae1')
    struct.pack_into('<HHHH', header, 24, 0x3e, 3, 0xfffe, 9)
    struct.pack_into('<H', header, 32, 6)
    struct.pack_into('<IIIIIIIII', header, 40, 0, 1, 0, 0, 4096, END, 0, END, 0)
    struct.pack_into('<109I', header, 76, fat_sector, *([FREE]*108))
    return bytes(header)+b''.join(sectors)+struct.pack('<128I', *fat)

cases = []
def add(name, data, *, expected=None, reason=''):
    if isinstance(data, str): data = data.encode()
    result = service.extract(file_content=data, filename=name)
    row = dict(name=name, content=base64.b64encode(data).decode(), text=result.text,
               tables=len(result.tables), pages=result.page_count, error=bool(result.error))
    if expected is not None: row.update(improved_text=expected, improvement=reason)
    cases.append(row)

for encoding in ['utf-8', 'utf-8-sig', 'utf-16', 'utf-16be', 'cp1252', 'latin1']:
    text = 'Notice: café — payment due.\nSecond line.' if encoding != 'latin1' else 'Café\x81 notice.'
    data = text.encode(encoding)
    if encoding == 'utf-16be': data = b'\xfe\xff'+data
    add('notice-'+encoding+'.txt', data)
for suffix in ['md','markdown','json']:
    add('plain.'+suffix, '{"notice":"Payment due"}' if suffix=='json' else '# Terms\nPayment due in 30 days.')
for i, raw in enumerate(['A,B\n1,2\n', 'A,B,C\n1,2\n', '"Party\nname",Term\nAcme,Net 30\n', 'A,B\n"x,y",\n']):
    add(f'table-{i}.csv',raw)
add('table.tsv','Party\tTerm\nAcme\tNet 30\n')
add('literal-pipe.csv','Party,Term\nAcme,A|B\n', expected='| Party | Term |\n| --- | --- |\n| Acme | A\\|B |',reason='Escape literal pipe characters so a cell cannot create extra Markdown columns.')
add('body.html','<head><title>Hidden</title></head><body><p>Terms &amp; conditions</p><div>Payment <b>due</b>.</div><script>hidden()</script></body>')
add('meta.html','<head><meta charset="UTF-8"><link rel="stylesheet" href="x"></head><body><p>Visible agreement</p></body>',expected='Visible agreement',reason='Void meta/link tags must not leave the HTML skip-depth counter stuck.')
add('spaces.html','<p><b>Pay</b> <i>now</i></p>', expected='Pay now', reason='Preserve whitespace between inline elements instead of joining separate words.')
for i, raw in enumerate([r'{\rtf1\ansi Rent is \b due\b0 .\par Second paragraph.}',r'{\rtf1\ansi\ansicpg1252 It\u8217?s due \'97 now.}',r'{\rtf1\ansi {\fonttbl{\f0 Arial;}}A {\field{\*\fldinst HYPERLINK "https://example.test"}{\fldrslt link}}.}',r'{\rtf1\ansi Emoji: \u-10179?\u-8704?}',r'{\rtf1\ansi Name\tab Term\par Acme\tab 30 days}']):
    add(f'document-{i}.rtf',raw)
for codepage, encoding in [(932,'cp932'),(65001,'utf-8')]:
    encoded='契約'.encode(encoding)
    add(f'codepage-{codepage}.rtf',r'{\rtf1\ansi\ansicpg'+str(codepage)+' '+''.join("\\'%02x"%b for b in encoded)+'.}')
for codepage, encoding, text in [(874,'cp874','สัญญา'),(852,'cp852','Příloha'),(866,'cp866','Договор'),(21866,'koi8_u','Угода'),(28592,'iso8859_2','Příloha')]:
    add(f'codepage-{codepage}.rtf',r'{\rtf1\ansi\ansicpg'+str(codepage)+' '+''.join("\\'%02x"%b for b in text.encode(encoding))+'.}',expected=text+'.',reason='Decode the declared code page, including ISO/KOI8 identifiers that are not Python cp aliases, instead of falling back to raw RTF.')
for charset,codec,text in [(204,'cp1251','Договор'),(128,'cp932','契約'),(238,'cp1250','Příloha')]:
    escaped=''.join("\\'%02x"%b for b in text.encode(codec))
    add(f'font-charset-{charset}.rtf',r'{\rtf1\ansi\deff0{\fonttbl{\f0\fnil\fcharset0 Arial;}{\f1\fnil\fcharset'+str(charset)+r' Arial;}}\f0 Agreement: {\f1 '+escaped+r'}\f0 signed.}')
add('message.eml','From: Counsel <counsel@example.test>\nTo: Client <client@example.test>\nSubject: =?UTF-8?B?Q2Fmw6kgdGVybXM=?=\nMIME-Version: 1.0\nContent-Type: multipart/mixed; boundary="m"\n\n--m\nContent-Type: multipart/alternative; boundary="a"\n\n--a\nContent-Type: text/plain; charset=utf-8\nContent-Transfer-Encoding: quoted-printable\n\nPayment is due in 30 days.\n--a\nContent-Type: text/html\n\n<p>HTML alternative</p>\n--a--\n--m\nContent-Type: application/pdf\nContent-Disposition: attachment; filename="agreement.pdf"\nContent-Transfer-Encoding: base64\n\nYnl0ZXM=\n--m--\n')
add('html-mail.eml','From: a@example.test\nSubject: Terms\nContent-Type: text/html; charset=utf-8\n\n<p>Payment <b>due</b>.</p>')
odt = archive({'mimetype':'application/vnd.oasis.opendocument.text','content.xml':'<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text><text:h>Terms</text:h><text:p>Payment <text:span>due</text:span> today.</text:p></office:text></office:body></office:document-content>'})
add('agreement.odt',odt)

from openpyxl import Workbook
from datetime import datetime
from docx import Document
document=Document(); document.add_heading('Document format acceptance',level=1); document.add_paragraph('Payment is due in thirty days. Confirm receipt in writing.')
table=document.add_table(rows=2,cols=2)
for r,row in enumerate([['Party','Notice'],['Acme','30 days']]):
    for c,value in enumerate(row): table.cell(r,c).text=value
document.add_paragraph('End of agreement.')
buffer=io.BytesIO(); document.save(buffer); add('agreement.docx',normalize_zip(buffer.getvalue()))
wb=Workbook(); wb.properties.created=wb.properties.modified=datetime(2000,1,1)
ws=wb.active; ws.title='Invoices'; ws.append(['Party','Amount']); ws.append(['Acme',1250]); ws.append(['Beta',40.5]); ws=wb.create_sheet('Terms'); ws.append(['Notice','Days']); ws.append(['Written',30])
buffer=io.BytesIO(); wb.save(buffer); add('workbook.xlsx',normalize_zip(buffer.getvalue()))
import xlwt
wb=xlwt.Workbook()
for name, rows in [('Invoices',[['Party','Amount'],['Acme',1250],['Beta',40.5]]),('Terms',[['Notice','Days'],['Written',30]])]:
    ws=wb.add_sheet(name)
    for r,row in enumerate(rows):
        for c,value in enumerate(row): ws.write(r,c,value)
buffer=io.BytesIO(); wb.save(buffer)
add('workbook.xls',buffer.getvalue(),expected='## Invoices\n\n| Party | Amount |\n| --- | --- |\n| Acme | 1250 |\n| Beta | 40.5 |\n\n## Terms\n\n| Notice | Days |\n| --- | --- |\n| Written | 30 |',reason='Whole-number Excel values use their displayed number rather than an artificial .0 suffix.')

# Independent byte-level fixtures for the pre-1997 formats xlrd accepts. The
# expected values come from the pinned Python service and xlrd, never the Go
# reader. Format layout reference: https://www.openoffice.org/sc/excelfileformat.pdf
import xlrd
def record(kind, body=b''): return struct.pack('<HH',kind,len(body))+body
def bof(version, kind=0x10):
    if version == 2: return record(0x9,struct.pack('<HH',2,kind))
    if version in (3,4): return record({3:0x209,4:0x409}[version],struct.pack('<HHH',version<<8,kind,0))
    return record(0x809,struct.pack('<HHHH',0x500,kind,2412,1993))
def xf(version): return b'' if version==2 else record({3:0x243,4:0x443,5:0xe0}[version],bytes(16 if version==5 else 12))
def cell(version,row,col): return struct.pack('<HH',row,col)+bytes(3 if version==2 else 2)
def label(version,row,col,value,codec='cp1252'):
    data=value.encode(codec)
    return record(4 if version==2 else 0x204,cell(version,row,col)+struct.pack('<B' if version==2 else '<H',len(data))+data)
def legacy_sheet(version,codepage=1252,codec='cp1252',label_text='Café',extra=b''):
    return (bof(version)+record(0x42,struct.pack('<H',codepage))+xf(version)
        +label(version,0,0,'Party')+label(version,0,1,'Amount')
        +label(version,1,0,label_text,codec)+record(3 if version==2 else 0x203,cell(version,1,1)+struct.pack('<d',40.5))
        +extra+record(0xa))
def legacy_book(sheets,codepage=1252):
    start=bof(5,5)+record(0x42,struct.pack('<H',codepage))+xf(5)
    names=[name.encode('cp1252') for name,_ in sheets]
    offset=len(start)+sum(4+7+len(name) for name in names)+4
    bounds=[]
    for name,(_,data) in zip(names,sheets):
        bounds.append(record(0x85,struct.pack('<IBBB',offset,0,0,len(name))+name)); offset+=len(data)
    return start+b''.join(bounds)+record(0xa)+b''.join(data for _,data in sheets)
def add_legacy(name,data):
    wb=xlrd.open_workbook(file_contents=data,logfile=io.StringIO())
    parts=[]
    for ws in wb.sheets():
        rows=[]
        for row in ws.get_rows():
            cells=[]
            for c in row:
                if c.ctype==xlrd.XL_CELL_BOOLEAN: value='TRUE' if c.value else 'FALSE'
                elif c.ctype==xlrd.XL_CELL_ERROR: value=xlrd.error_text_from_code[c.value]
                elif c.ctype==xlrd.XL_CELL_NUMBER: value=str(int(c.value)) if c.value.is_integer() else str(c.value)
                else: value=str(c.value).strip()
                cells.append(value)
            if any(cells): rows.append(cells)
        if rows: parts.append(f'## {ws.name}\n\n'+service._rows_to_markdown(rows))
    add(name,data,expected='\n\n'.join(parts),reason='Keep all independently decoded cells; display whole numbers, booleans and Excel errors as readable values instead of numeric type codes.')
    if cases[-1]['error']: raise ValueError(f'Upstream rejected {name}')
for version in (2,3,4): add_legacy(f'biff{version}.xls',legacy_sheet(version))
add_legacy('biff2-notice.xls',legacy_sheet(2,label_text='Notice is 30 days.'))
for cp,codec,text in [(1251,'cp1251','Договор'),(932,'cp932','契約'),(10000,'mac_roman','Café')]:
    add_legacy(f'biff5-codepage-{cp}.xls',compound({'Book':legacy_book([('Terms',legacy_sheet(5,cp,codec,text))],cp)}))
extra=(label(5,2,0,'Boolean')+record(0x205,cell(5,2,1)+bytes([1,0]))
    +label(5,3,0,'Error')+record(0x205,cell(5,3,1)+bytes([7,1]))
    +label(5,4,0,'Scaled number')+record(0x27e,cell(5,4,1)+struct.pack('<I',(12345<<2)|3))
    +record(0xbd,struct.pack('<HH',5,0)+struct.pack('<HIHIH',0,(-2<<2 & 0xffffffff)|2,0,(30<<2)|2,1))
    +label(5,6,0,'Cached formula')+record(6,cell(5,6,1)+struct.pack('<d',12.5)+bytes(8))
    +label(5,7,0,'Cached text')+record(6,cell(5,7,1)+bytes([0,0,0,0,0,0,255,255])+bytes(8))
    +record(0x207,struct.pack('<H',300)+b'A'*100)+record(0x3c,b'A'*200)
    +bof(5,0x20)+record(0xa)+label(5,8,0,'After embedded chart')+label(5,8,1,'Preserved'))
book=legacy_book([('First',legacy_sheet(5,extra=extra)),('Second',legacy_sheet(5,label_text='Other'))])
add_legacy('biff5-cached-values.xls',compound({'Book':book}))
add_legacy('biff5-raw.xls',book)
names=[b'First',b'Second']
prefix=bof(4,0x100)+record(0x42,struct.pack('<H',1252))
bounds=b''.join(record(0x85,bytes([len(n)])+n) for n in names)
header=prefix+record(0x8e,struct.pack('<I',len(prefix)+8+len(bounds)))+bounds
body=b''
for n in names:
    sheet=legacy_sheet(4)
    body+=record(0x8f,struct.pack('<I',len(sheet))+bytes([len(n)])+n)+sheet
add_legacy('biff4-workbook.xls',header+body+record(0xa))
from pptx import Presentation
from pptx.util import Inches
prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text='Terms'
table=slide.shapes.add_table(2,2, Inches(1), Inches(2), Inches(6), Inches(2)).table
for r,row in enumerate([['Party','Term'],['Acme','30 days']]):
    for c,value in enumerate(row): table.cell(r,c).text=value
slide.notes_slide.notes_text_frame.text='Confirm receipt.'
slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text='Next steps'; slide.placeholders[1].text='Send notice.'
buffer=io.BytesIO(); prs.save(buffer); add('presentation.pptx',normalize_zip(buffer.getvalue()))

text='Rent is due — 契約 applies.\rParty\x07Term\x07\x07Acme\x0730 days\x07\x07'
encoded=text.encode('utf-16le'); word=bytearray(1024)+encoded
struct.pack_into('<H',word,0,0xa5ec); struct.pack_into('<I',word,0x4c,len(encoded)//2)
plc=struct.pack('<II',0,len(encoded)//2)+b'\0\0'+struct.pack('<I',1024)+b'\0\0'
clx=b'\x02'+struct.pack('<I',len(plc))+plc
struct.pack_into('<II',word,0x1a2,0,len(clx))
add('legacy.doc',compound({'WordDocument':word,'0Table':clx}))
def prop(tag,value): return ('__substg1.0_'+tag+'001F',(value+'\0').encode('utf-16le'))
streams=dict([prop('0C1A','Counsel'),prop('5D01','counsel@example.test'),prop('0E04','Client'),prop('0037','Terms'),prop('1000','Payment is due in 30 days.')])
name,value=prop('3707','agreement.pdf'); streams['__attach_version1.0_#00000000/'+name]=value
properties=bytearray(48); struct.pack_into('<IIQ',properties,32,0x00390040,0,133485408000000000); streams['__properties_version1.0']=properties
add('message.msg',compound(streams))

for name,data in [('bad.doc',b'not word'),('bad.xlsx',b'not zip'),('bad.pptx',b'not zip'),('binary.txt',b'A\0B'),('empty.txt',b'')]: add(name,data)

raw=source('backend/app/services/courts.py')
courts=next(ast.literal_eval(n.value) for n in ast.parse(raw).body if isinstance(n,ast.AnnAssign) and n.target.id=='COURTS')
(destination/'courts.json').write_text(json.dumps(courts,indent=2)+'\n')
(destination/'testdata/text_extraction.json').write_text(json.dumps(dict(baseline=BASELINE,sources=sources,cases=cases),indent=2,ensure_ascii=True)+'\n')
print(f'Wrote {len(cases)} extraction cases and {len(courts)} jurisdictions')
