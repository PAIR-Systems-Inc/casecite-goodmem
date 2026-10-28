"""Export a review table: an .xlsx with a cited Results sheet, or a flat CSV.

Results holds one row per document and one column per field. Every answered
cell links to its row on the Citations sheet, which carries the quote, its
offsets and whether it verified, so a figure in the spreadsheet can always be
traced back to the words in the document that support it.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from app.models.review_tables import CELL_ERROR, CELL_EXTRACTED, CELL_NOT_FOUND, CELL_UNVERIFIED

RESULTS_SHEET = "Results"
CITATIONS_SHEET = "Citations"
CITATION_HEADERS = (
    "Document",
    "Field",
    "Value",
    "Quote",
    "Span start",
    "Span end",
    "Verified",
    "Status",
    "Error",
)

_STATUS_LABEL = {
    CELL_EXTRACTED: "verified",
    CELL_UNVERIFIED: "unverified",
    CELL_NOT_FOUND: "not found",
    CELL_ERROR: "error",
}


def _display_value(cell: dict[str, Any] | None) -> str:
    if cell is None:
        return ""
    status = cell.get("status")
    if status == CELL_NOT_FOUND:
        return "—"
    if status == CELL_ERROR:
        return "error"
    return str(cell.get("value") or "")


def _cell_index(cells: list[dict[str, Any]]) -> dict[tuple[str, str], dict[str, Any]]:
    return {(c["document_id"], c["field_key"]): c for c in cells}


def render_xlsx(
    name: str,
    fields: list[dict[str, Any]],
    documents: list[dict[str, Any]],
    cells: list[dict[str, Any]],
) -> bytes:
    """Workbook bytes: Results (hyperlinked to) Citations."""
    wb = Workbook()
    results = wb.active
    assert results is not None
    results.title = RESULTS_SHEET
    citations = wb.create_sheet(CITATIONS_SHEET)

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8ECF2")
    link_font = Font(color="1F4E99", underline="single")
    unverified_fill = PatternFill("solid", fgColor="FFF4D6")
    error_fill = PatternFill("solid", fgColor="FBE3E3")

    for col, title in enumerate(["Document", *[f["name"] for f in fields]], start=1):
        c = results.cell(row=1, column=col, value=title)
        c.font = header_font
        c.fill = header_fill
    for col, title in enumerate(CITATION_HEADERS, start=1):
        c = citations.cell(row=1, column=col, value=title)
        c.font = header_font
        c.fill = header_fill

    by_key = _cell_index(cells)
    citation_row = 2
    for r, doc in enumerate(documents, start=2):
        results.cell(row=r, column=1, value=doc["filename"])
        for col, field in enumerate(fields, start=2):
            cell = by_key.get((doc["id"], field["key"]))
            out = results.cell(row=r, column=col, value=_display_value(cell))
            if cell is None or cell.get("status") == CELL_NOT_FOUND:
                continue
            citations.append(
                [
                    doc["filename"],
                    field["name"],
                    cell.get("value") or "",
                    cell.get("quote") or "",
                    cell.get("span_start"),
                    cell.get("span_end"),
                    "yes" if cell.get("verified") else "no",
                    _STATUS_LABEL.get(str(cell.get("status")), str(cell.get("status"))),
                    cell.get("error") or "",
                ]
            )
            out.hyperlink = Hyperlink(
                ref=out.coordinate,
                location=f"'{CITATIONS_SHEET}'!A{citation_row}",
                display=out.value,
            )
            out.font = link_font
            if cell.get("status") == CELL_UNVERIFIED:
                out.fill = unverified_fill
            elif cell.get("status") == CELL_ERROR:
                out.fill = error_fill
            citation_row += 1

    for ws, widths in (
        (results, [36] + [28] * len(fields)),
        (citations, [36, 24, 28, 60, 10, 10, 9, 11, 30]),
    ):
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "B2" if ws is results else "A2"
    for row in citations.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.properties.title = name
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_csv(
    fields: list[dict[str, Any]],
    documents: list[dict[str, Any]],
    cells: list[dict[str, Any]],
) -> str:
    """Flat CSV: one row per (document, field) with value, quote and verification."""
    by_key = _cell_index(cells)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(CITATION_HEADERS)
    for doc in documents:
        for field in fields:
            cell = by_key.get((doc["id"], field["key"]))
            if cell is None:
                writer.writerow([doc["filename"], field["name"], "", "", "", "", "", "", ""])
                continue
            writer.writerow(
                [
                    doc["filename"],
                    field["name"],
                    cell.get("value") or "",
                    cell.get("quote") or "",
                    cell.get("span_start") if cell.get("span_start") is not None else "",
                    cell.get("span_end") if cell.get("span_end") is not None else "",
                    "yes" if cell.get("verified") else "no",
                    _STATUS_LABEL.get(str(cell.get("status")), str(cell.get("status"))),
                    cell.get("error") or "",
                ]
            )
    return buf.getvalue()
